"""Bounded XLSX interchange for the single Dataset contract; never evaluate cells."""
from datetime import date, datetime
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from time import perf_counter
from zipfile import ZipFile, BadZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from pydantic import ValidationError
from backend.app import contracts as c
from backend.app.data.validation import validate_dataset
from backend.app.planning.inputs import planning_input_failures

MAX_FILE = 16 * 1024 * 1024
MAX_EXPANDED = 160 * 1024 * 1024
MAX_ENTRY = 120 * 1024 * 1024
MAX_ENTRIES = 100
MAX_RATIO = 250
MAX_ROWS = 170000
MAX_CELL = 4000
TABLES = {
    'Products': ('products', c.Product, 60, ('sku',)),
    'Locations': ('locations', c.Location, 5, ('location_id',)),
    'Assortment': ('assortment', c.Assortment, 240, ('sku', 'location_id')),
    'DemandHistory': ('demand_history', c.Observation, 100800, ('sku', 'location_id', 'day')),
    'Inventory': ('inventory', c.Inventory, 300, ('sku', 'location_id')),
    'OpenOrders': ('open_orders', c.OpenOrder, 5000, ('external_id',)),
    'OpenTransfers': ('open_transfers', c.OpenTransfer, 5000, ('external_id',)),
    'SupplierOffers': ('supplier_offers', c.SupplierOffer, 720, ('offer_id',)),
    'Suppliers': ('suppliers', c.Supplier, 12, ('supplier_id',)),
    'SupplierCapacity': ('supplier_capacity', c.SupplierCapacity, 40320, ('supplier_id', 'sku', 'dispatch_date')),
    'TransferLanes': ('transfer_lanes', c.TransferLane, 20, ('source', 'destination')),
    'Payables': ('payables', c.Payable, 10000, ('external_id',)),
    'Budgets': ('budgets', c.Budget, 30, ('week_start',)),
    'Events': ('events', c.Event, 100, ('event_id',)),
}
SETTINGS = ['schema_version', 'dataset_id', 'synthetic', 'declared_empty', *c.Settings.model_fields]
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class WorkbookIssue(c.Contract):
    sheet: str
    row: int | None = None
    field: str | None = None
    severity: str = 'error'
    code: str
    guidance: str


class WorkbookError(ValueError):
    def __init__(self, issues):
        self.issues = issues
        super().__init__('Workbook validation failed. Correct the listed fields.')


def problem(code, guidance, sheet='Workbook', row=None, field=None):
    return WorkbookIssue(sheet=sheet, row=row, field=field, code=code, guidance=guidance)


def append(ws, values):
    # Literal strings, including user notes beginning with '=', remain text in exports.
    cells = []
    for value in values:
        if isinstance(value, (list, dict)):
            value = json.dumps(value, separators=(',', ':'), ensure_ascii=False)
        if isinstance(value, (date, datetime)):
            value = value.isoformat()
        if isinstance(value, float):
            # Excel numeric cells truncate binary-float round-trip precision.
            value = repr(value)
        cell = WriteOnlyCell(ws, value=value)
        if isinstance(value, str):
            cell.data_type = 's'
        cells.append(cell)
    ws.append(cells)


def template(data=None):
    wb = Workbook(write_only=True)
    for name, (field, model, _, _) in TABLES.items():
        ws = wb.create_sheet(name)
        append(ws, list(model.model_fields))
        for record in getattr(data, field, []) if data else []:
            append(ws, list(record.model_dump(mode='json').values()))
    ws = wb.create_sheet('Settings'); append(ws, ['key', 'value'])
    settings = data.settings.model_dump(mode='json') if data else {}
    metadata = {k: getattr(data, k) for k in SETTINGS[:4]} if data else {}
    for key in SETTINGS:
        append(ws, [key, metadata.get(key, settings.get(key))])
    ws = wb.create_sheet('Reconciliation')
    append(ws, ['external_id', 'confirmed_units', 'executed_units'])
    ws = wb.create_sheet('Instructions')
    append(ws, ['sheet', 'field', 'required', 'definition'])
    append(ws, ['ALL', 'units', True, 'Base units only; money SAR per base unit. ISO YYYY-MM-DD dates; timestamps include offset. Lists/maps are JSON. No formulas.'])
    append(ws, ['Inventory', 'reserved', True, 'Committed outside this demand forecast; deducted once with blocked stock.'])
    append(ws, ['Budgets', 'new_commitment_cap', True, 'Remaining authority for NEW orders; existing obligations consume payment_ceiling, not this cap.'])
    append(ws, ['SupplierCapacity', 'available_units', True, 'Remaining dated availability for NEW orders; zero is explicit, missing is unknown.'])
    append(ws, ['Settings', 'declared_empty', True, 'JSON list of intentionally empty open_orders, open_transfers, payables. Missing sheets are never allowed.'])
    for name, (_, model, _, keys) in TABLES.items():
        schema = model.model_json_schema()['properties']
        for field, info in model.model_fields.items():
            append(ws, [name, field, info.is_required(), ('PRIMARY KEY; ' if field in keys else '') + json.dumps(schema[field], separators=(',', ':'))])
    for field, info in c.Settings.model_fields.items():
        append(ws, ['Settings', field, info.is_required(), json.dumps(c.Settings.model_json_schema()['properties'][field])])
    ws = wb.create_sheet('Examples')
    append(ws, ['NOT IMPORTED', 'field', 'example'])
    append(ws, ['Settings', 'declared_empty', '["open_transfers"]'])
    append(ws, ['Settings', 'class_targets', '{"A":0.98,"B":0.95,"C":0.9}'])
    append(ws, ['Locations', 'open_weekdays', '[0,1,2,3,4,5,6]'])
    out = BytesIO(); wb.save(out); return out.getvalue()


def preflight(content):
    if len(content) > MAX_FILE:
        raise WorkbookError([problem('file_limit', 'Use an XLSX file of at most 16 MiB.')])
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ENTRIES or len({e.filename for e in entries}) != len(entries):
                raise WorkbookError([problem('zip_entries', 'Workbook has too many or duplicate ZIP entries.')])
            if sum(e.file_size for e in entries) > MAX_EXPANDED:
                raise WorkbookError([problem('zip_expanded_limit', 'Expanded workbook exceeds 160 MiB.')])
            for entry in entries:
                metadata_limit = 8*1024*1024 if entry.filename == 'xl/sharedStrings.xml' else 1024*1024 if entry.filename == 'xl/styles.xml' else MAX_ENTRY
                if entry.file_size > metadata_limit:
                    raise WorkbookError([problem('zip_metadata_limit', 'Workbook metadata exceeds safe parsing limits.')])
                if entry.flag_bits & 1 or entry.file_size > MAX_ENTRY or entry.file_size / max(1, entry.compress_size) > MAX_RATIO:
                    raise WorkbookError([problem('zip_entry_limit', 'Encrypted, oversized or excessively compressed ZIP entry; save a plain XLSX workbook.')])
                if any(x in entry.filename.lower() for x in ('vbaproject', 'externallinks', 'embeddings/')):
                    raise WorkbookError([problem('active_content', 'Remove macros, external workbook links and embedded objects.')])
            if '[Content_Types].xml' not in archive.namelist() or 'xl/workbook.xml' not in archive.namelist():
                raise WorkbookError([problem('xlsx_content', 'File must contain a valid XLSX workbook.')])
            # Stream all XML with entity protection before openpyxl allocates workbook objects.
            for entry in entries:
                if entry.filename.endswith(('.xml', '.rels')):
                    with archive.open(entry) as stream:
                        for _, node in iterparse(stream, events=('end',)):
                            if 'macroEnabled' in str(node.attrib.get('ContentType', '')):
                                raise WorkbookError([problem('macro_content', 'Macro-enabled workbooks are unsupported.')])
                            if node.tag.endswith('}Relationship') and node.attrib.get('TargetMode') == 'External':
                                raise WorkbookError([problem('external_link', 'Remove external relationships from the workbook.')])
                            node.clear()
            return sum(e.file_size for e in entries)
    except WorkbookError:
        raise
    except Exception as error:
        raise WorkbookError([problem('malformed_xlsx', 'File is malformed, encrypted or unsafe XML. Save an ordinary unencrypted XLSX workbook.')]) from error


def cell_value(cell):
    value = cell.value
    if cell.data_type in ('f', 'e'):
        raise ValueError('formula')
    if isinstance(value, str):
        if len(value) > MAX_CELL:
            raise ValueError('cell_limit')
        if value.startswith(('[', '{')):
            return json.loads(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10] if isinstance(value, datetime) and value.time().isoformat() == '00:00:00' else value.isoformat()
    return value


def parse(content, filename='data.xlsx'):
    started = perf_counter()
    if Path(filename).suffix.lower() != '.xlsx':
        raise WorkbookError([problem('extension', 'Only .xlsx is supported; .xls and .xlsm are rejected.')])
    expanded = preflight(content)
    # Isolated file exists only within this request, including unexpected exceptions.
    with TemporaryDirectory(prefix='planning-import-') as directory:
        path = Path(directory) / 'input.xlsx'; path.write_bytes(content)
        try:
            wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                data, issues, counts, executions = _parse(wb)
            finally:
                wb.close()
        except WorkbookError:
            raise
        except Exception as error:
            raise WorkbookError([problem('malformed_xlsx', 'Workbook cannot be read safely. Save a new XLSX using the template.')]) from error
    return data, issues, dict(source_bytes=len(content), expanded_bytes=expanded, sheet_count=len(counts), rows=counts,
                             parse_validation_ms=round((perf_counter()-started)*1000, 2), temporary_cleanup=True, reconciliation=executions)


def _parse(wb):
    required = set(TABLES) | {'Settings'}
    extra = set(wb.sheetnames) - required - {'Instructions', 'Examples', 'Reconciliation'}
    issues = [problem('missing_sheet', 'Restore the required named sheet.', name) for name in sorted(required-set(wb.sheetnames))]
    issues += [problem('unexpected_sheet', 'Use only template operational sheets, Instructions and Examples.', name) for name in sorted(extra)]
    if issues: raise WorkbookError(issues)
    raw = {}; counts = {}; positions = {}; total = 0
    for name in sorted(required):
        ws = wb[name]
        if ws.max_column and ws.max_column > 40 or ws.max_row and ws.max_row > MAX_ROWS:
            raise WorkbookError([problem('sheet_limit', 'Sheet dimensions exceed the supported limits.', name)])
        # Do not trust a forged worksheet dimension to hide rows or formulas.
        ws.reset_dimensions()
        rows = ws.iter_rows()
        headers = [cell.value for cell in next(rows, [])]
        expected = list(TABLES[name][1].model_fields) if name in TABLES else ['key', 'value']
        if len(headers) != len(expected) or set(headers) != set(expected):
            issues.append(problem('headers', 'Provide each template field exactly once; unknown fields are unsupported.', name, 1)); continue
        values = []; seen = set(); counts[name] = 0; positions[name] = []
        limit = TABLES[name][2] if name in TABLES else len(SETTINGS)
        for row_number, cells in enumerate(rows, 2):
            if row_number > limit+1:
                raise WorkbookError([problem('row_limit', f'Sheet exceeds its {limit}-row limit.', name, row_number)])
            if len(cells) > len(headers):
                issues.append(problem('extra_cells', 'Remove cells outside the defined columns.', name, row_number)); continue
            if all(cell.value is None for cell in cells): continue
            counts[name] += 1; total += 1
            if total > MAX_ROWS: raise WorkbookError([problem('total_rows', 'Workbook exceeds 170000 operational rows.')])
            record = {}
            for field, cell in zip(headers, cells):
                try: record[field] = cell_value(cell)
                except (ValueError, TypeError): issues.append(problem('cell_value', 'Use a literal value, no formulas/errors; JSON lists/maps must be valid and cells at most 4000 characters.', name, row_number, field))
            if name == 'Settings':
                key = record.get('key')
                if key not in SETTINGS or key in seen: issues.append(problem('setting_key', 'Use each documented Settings key exactly once.', name, row_number, 'key'))
                seen.add(key); values.append(record); continue
            field, model, _, keys = TABLES[name]
            try:
                row = model.model_validate(record)
                key = tuple(getattr(row, k) for k in keys)
                if key in seen: issues.append(problem('duplicate_key', 'Remove the duplicate primary key; do not merge conflicting records.', name, row_number, ','.join(keys)))
                seen.add(key); values.append(row); positions[name].append(row_number)
            except ValidationError as error:
                for e in error.errors(include_input=False):
                    issues.append(problem('field_'+e['type'], 'Correct the value using the field definition in Instructions.', name, row_number, '.'.join(map(str,e['loc']))))
            if len(issues) >= 100: raise WorkbookError(issues[:100])
        if name == 'Settings':
            settings = {r['key']: r.get('value') for r in values if r.get('key') in SETTINGS}
            for key in set(SETTINGS)-set(settings): issues.append(problem('missing_setting', 'Restore the Settings key and its explicit value.', name, None, key))
            raw.update({k: settings.get(k) for k in SETTINGS[:4]})
            raw['settings'] = {k: settings.get(k) for k in c.Settings.model_fields}
        else: raw[TABLES[name][0]] = values
    if issues: raise WorkbookError(issues[:100])
    try: data = c.Dataset.model_validate(raw)
    except ValidationError as error:
        raise WorkbookError([problem('dataset_'+e['type'], 'Complete the required dataset/settings values.', 'Settings', None, '.'.join(map(str,e['loc']))) for e in error.errors(include_input=False)[:100]]) from error
    mapping = {field: sheet for sheet, (field, *_rest) in TABLES.items()}
    references={'sku':{p.sku for p in data.products},'supplier_id':{s.supplier_id for s in data.suppliers},
                'location_id':{l.location_id for l in data.locations},'event_id':{e.event_id for e in data.events}}
    references['source']=references['destination']=references['location_id']
    detailed_codes=set()
    for sheet,(field,_,_,_) in TABLES.items():
        for row_number,record in zip(positions[sheet],getattr(data,field)):
            for key,known in references.items():
                value=getattr(record,key,None)
                if value is not None and value not in known:
                    code='unknown_sku' if key=='sku' else 'unknown_supplier' if key=='supplier_id' else 'unknown_event' if key=='event_id' else 'unknown_location'
                    detailed_codes.add(code);issues.append(problem(code,'Reference an identifier defined in the corresponding master sheet.',sheet,row_number,key))
            if field=='inventory':
                if record.as_of!=data.settings.as_of:
                    detailed_codes.add('conflicting_snapshot');issues.append(problem('conflicting_snapshot','Use the single Settings as_of date for all inventory rows.',sheet,row_number,'as_of'))
                if record.on_hand<record.blocked+record.reserved:
                    detailed_codes.add('negative_usable_stock');issues.append(problem('negative_usable_stock','Blocked plus reserved cannot exceed on_hand.',sheet,row_number,'on_hand'))
    issue_sheets={'missing_history':'DemandHistory','censored_history':'DemandHistory','delayed_history':'DemandHistory','history_dates':'DemandHistory',
        'observation_time':'DemandHistory','funding_coverage':'Budgets','budget_week':'Budgets','existing_payment_breach':'Budgets',
        'pack_conversion':'SupplierOffers','offer_dates':'SupplierOffers','active_dates':'Products','range_dates':'Assortment',
        'transaction_dates':'OpenOrders','received_stock':'OpenOrders','overdue_receipt':'OpenOrders','unknown_payable_link':'Payables',
        'event_dates':'Events','event_scope':'Events','overlapping_events':'Events','unknown_capacity':'SupplierCapacity'}
    for issue in validate_dataset(data):
        if issue.code in detailed_codes:continue
        sheet = next((mapping[field] for field in mapping if issue.message.startswith(field+':')), issue_sheets.get(issue.code,'Settings'))
        issues.append(WorkbookIssue(sheet=sheet, severity=issue.severity, code=issue.code, guidance=issue.message))
    for failure in planning_input_failures(data):
        issues.append(problem(failure.code, failure.message, 'SupplierCapacity' if 'capacity' in failure.code else 'Settings'))
    if any(i.severity == 'error' for i in issues):
        # A long list of quality warnings must never hide the blocking error.
        raise WorkbookError(sorted(issues,key=lambda i:i.severity!='error')[:100])
    executions=[]
    if 'Reconciliation' in wb.sheetnames:
        from backend.app.data.reconciliation import Execution
        ws=wb['Reconciliation'];ws.reset_dimensions();rows=ws.iter_rows()
        headers=[cell.value for cell in next(rows,[])]
        if headers!=list(Execution.model_fields): raise WorkbookError([problem('headers','Use the three documented reconciliation columns.','Reconciliation',1)])
        seen=set()
        for row_number,cells in enumerate(rows,2):
            if row_number>10001: raise WorkbookError([problem('row_limit','At most 10000 reconciliation rows.','Reconciliation')])
            if all(cell.value is None for cell in cells): continue
            try:
                if len(cells)>3: raise ValueError('extra columns')
                row=Execution.model_validate({key:cell_value(cell) for key,cell in zip(headers,cells)})
                if row.external_id in seen: raise ValueError('duplicate external ID')
                seen.add(row.external_id);executions.append(row.model_dump())
            except (ValueError,TypeError):
                raise WorkbookError([problem('execution_row','Use a unique external ID and nonnegative integer confirmed/executed quantities. No formulas.','Reconciliation',row_number)])
        counts['Reconciliation']=len(executions)
    return data, issues, counts, executions
