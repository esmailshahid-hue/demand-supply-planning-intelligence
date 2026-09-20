from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from openpyxl import load_workbook
import openpyxl
from backend.app.data.sample import generate_sample
from backend.app.data.workbook import template, parse, preflight, WorkbookError, TABLES
from backend.app.data.storage import LocalStorage, ObjectUnavailable, ObjectReference, configuration


@pytest.fixture(scope='module')
def fixture(): return generate_sample()


def edit(content,change):
    wb=load_workbook(BytesIO(content));change(wb);out=BytesIO();wb.save(out);wb.close();return out.getvalue()


def test_blank_has_all_headers_and_separate_examples():
    content=template();wb=load_workbook(BytesIO(content),read_only=True)
    assert set(TABLES)|{'Settings','Instructions','Examples'} <= set(wb.sheetnames)
    for sheet,(_,model,_,_) in TABLES.items():
        assert list(next(wb[sheet].values))==list(model.model_fields)
    wb.close()
    with pytest.raises(WorkbookError):parse(content)


@pytest.mark.parametrize('size',['fixture','full'])
def test_complete_lossless_round_trip(size):
    data=generate_sample(size);content=template(data);restored,issues,metrics=parse(content)
    assert restored.model_dump_json()==data.model_dump_json()
    assert metrics['temporary_cleanup'] and metrics['source_bytes']==len(content)
    assert metrics['rows']['DemandHistory']==len(data.demand_history)
    assert not any(i.severity=='error' for i in issues)


@pytest.mark.parametrize('change,code',[
    (lambda w:w.remove(w['Products']),'missing_sheet'),
    (lambda w:w.create_sheet('Unexpected'),'unexpected_sheet'),
    (lambda w:setattr(w['Products']['A1'],'value','unknown'),'headers'),
    (lambda w:setattr(w['Inventory']['D2'],'value',-1),'field_greater_than_equal'),
    (lambda w:setattr(w['Inventory']['D2'],'value','=1+1'),'cell_value'),
    (lambda w:setattr(w['Inventory']['A2'],'value','UNKNOWN'),'unknown_sku'),
    (lambda w:setattr(w['Inventory']['C2'],'value','2020-01-01'),'conflicting_snapshot'),
    (lambda w:setattr(w['Settings']['B5'],'value','[]'),'undeclared_empty'),
    (lambda w:setattr(w['SupplierOffers']['I2'],'value',15),'field_less_than_equal'),
    (lambda w:setattr(w['Products']['C2'],'value','case'),'field_literal_error'),
    (lambda w:w['Products'].append([c.value for c in w['Products'][2]]),'duplicate_key'),
])
def test_grouped_invalid_inputs(fixture,change,code):
    with pytest.raises(WorkbookError) as e:parse(edit(template(fixture),change))
    assert code in {i.code for i in e.value.issues}
    assert all(i.sheet and i.guidance and i.severity in ('error','warning') for i in e.value.issues)
    assert any(i.severity=='error' for i in e.value.issues)


@pytest.mark.parametrize('name,content',[('test.xls',b'bad'),('test.xlsm',b'bad'),('test.xlsx',b'bad')])
def test_unsupported_and_malformed(name,content):
    with pytest.raises(WorkbookError):parse(content,name)


def test_limits_before_loading(fixture,monkeypatch):
    import backend.app.data.workbook as module
    content=template(fixture)
    monkeypatch.setattr(module,'MAX_FILE',len(content)-1)
    with pytest.raises(WorkbookError,match='validation'):parse(content)
    monkeypatch.setattr(module,'MAX_FILE',16*1024*1024)
    monkeypatch.setattr(module,'MAX_EXPANDED',1000)
    with pytest.raises(WorkbookError) as error:preflight(content)
    assert error.value.issues[0].code=='zip_expanded_limit'


def test_zip_bomb_and_macro(fixture):
    out=BytesIO()
    with ZipFile(out,'w',ZIP_DEFLATED) as z:z.writestr('bomb',b'0'*1000000)
    with pytest.raises(WorkbookError) as error:preflight(out.getvalue())
    assert error.value.issues[0].code=='zip_entry_limit'
    content=BytesIO(template(fixture))
    with ZipFile(content,'a') as z:z.writestr('xl/vbaProject.bin',b'macro')
    with pytest.raises(WorkbookError) as error:preflight(content.getvalue())
    assert error.value.issues[0].code=='active_content'


def test_cleanup_on_unexpected_parse_exception(fixture,monkeypatch):
    import backend.app.data.workbook as module
    captured=[]
    def broken(path,**kwargs):captured.append(Path(path).parent);raise RuntimeError('secret path')
    monkeypatch.setattr(openpyxl,'load_workbook',broken)
    with pytest.raises(WorkbookError) as error:parse(template(fixture))
    assert 'secret' not in str(error.value)
    assert captured and not captured[0].exists()


def test_private_storage_ownership_expiry_cleanup_and_host_block(monkeypatch):
    clock=[0];store=LocalStorage(ttl=10,clock=lambda:clock[0]);owner='a'*43;other='b'*43
    try:
        ref=store.put(owner,b'private','dataset')
        assert store.get(owner,ref,'dataset')==b'private'
        with pytest.raises(ObjectUnavailable):store.get(other,ref,'dataset')
        with pytest.raises(ObjectUnavailable):store.get(owner,ref,'upload')
        clock[0]=11;store.sweep();assert not list(store.root.iterdir())
        with pytest.raises(ObjectUnavailable):store.get(owner,ref,'dataset')
        ref=store.put(owner,b'next','dataset');store.reset(owner);assert not list(store.root.iterdir())
    finally:store.close()
    monkeypatch.setenv('VERCEL','1');monkeypatch.setenv('PLANNING_UPLOAD_STORAGE','local')
    assert not configuration()['enabled']


def test_operational_row_limit(fixture,monkeypatch):
    import backend.app.data.workbook as module
    monkeypatch.setitem(module.TABLES,'Products',('products',module.c.Product,5,('sku',)))
    with pytest.raises(WorkbookError) as error:parse(template(fixture))
    assert error.value.issues[0].code=='row_limit'


def test_reference_error_uses_physical_row_after_blank(fixture):
    def change(w):
        w['Inventory'].insert_rows(2)
        w['Inventory']['A3']='UNKNOWN'
    with pytest.raises(WorkbookError) as error:parse(edit(template(fixture),change))
    issue=next(i for i in error.value.issues if i.code=='unknown_sku')
    assert issue.sheet=='Inventory' and issue.row==3 and issue.field=='sku'


def test_warning_overflow_cannot_hide_blocking_error(fixture,monkeypatch):
    import backend.app.data.workbook as module
    from backend.app.contracts import Issue
    warnings=[Issue(severity='warning',code='quality',message='Review quality.') for _ in range(110)]
    monkeypatch.setattr(module,'validate_dataset',lambda data:warnings+[Issue(severity='error',code='funding_coverage',message='Missing coverage.')])
    with pytest.raises(WorkbookError) as error:parse(template(fixture))
    assert error.value.issues[0].code=='funding_coverage'


@pytest.mark.parametrize('invalid',[False,True])
def test_parse_files_deleted_after_success_or_validation_error(fixture,monkeypatch,invalid):
    import backend.app.data.workbook as module
    paths=[];original=openpyxl.load_workbook
    def capture(path,**kwargs):
        paths.append(Path(path).parent);return original(path,**kwargs)
    monkeypatch.setattr(openpyxl,'load_workbook',capture)
    content=template(fixture)
    if invalid:
        content=edit(content,lambda w:setattr(w['Inventory']['D2'],'value',-1))
        with pytest.raises(WorkbookError):parse(content)
    else:parse(content)
    assert paths and all(not path.exists() for path in paths)
